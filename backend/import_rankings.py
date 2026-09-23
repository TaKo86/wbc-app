"""
Import current WBC NZ rankings from the WBC_for_Regy.xlsx handoff spreadsheet.

Source-of-truth decisions (confirmed with WBC):
  - Men: the PRO sheet feeds Pro-level titles and the AMATEUR sheet feeds
    Amateur-level titles, independently of each other and of the
    "Consolidated" sheet's men's section (which mixes pro/amateur names and
    is not used here).
  - Women: the Consolidated sheet's "Woman" section is the source, and is
    duplicated into BOTH the Pro and Amateur title for each weight class,
    since that section has no separate pro/amateur split of its own.

This only touches Fighter / WeightClass / Title / Ranking rows, plus a single
placeholder TitleFight per newly-seen champion so the app's "champion = whoever
won the most recent title fight" logic has something to compute from. It never
overwrites real fight history: if a title already has at least one TitleFight
on file, the spreadsheet's Champion cell is reported as skipped rather than
applied automatically.

Rankings ARE replaced wholesale per title on each run (the sheet is treated as
the current state of the board), so this is safe to re-run after the
spreadsheet is updated.

Usage (after backend/db/schema.sql has been applied):
    python import_rankings.py path/to/WBC_for_Regy.xlsx
"""
import re
import sys
from datetime import date

import openpyxl

from app import models
from app.database import SessionLocal, engine, Base

NAME_BY_CODE = {
    "BW": "Bantamweight", "SBW": "Super-bantamweight", "FW": "Featherweight",
    "SFW": "Super-featherweight", "LW": "Lightweight", "SLW": "Super-lightweight",
    "WW": "Welterweight", "SWW": "Super-welterweight", "MW": "Middleweight",
    "SMW": "Super-middleweight", "LHW": "Light-heavyweight", "CW": "Cruiserweight",
    "HW": "Heavyweight",
}

HEADER_RE = re.compile(r"^([A-Za-z]+)\s+([+\d.]+)\s*kg", re.IGNORECASE)
FOOTNOTE_RE = re.compile(r"\*+\s*$")
WORLD_CHAMP_RE = re.compile(r"\bWorld Champ\b\.?\s*$", re.IGNORECASE)
REQ_BY_STARS = {
    1: "One MT win required",
    2: "Two MT wins required",
    3: "Thai experience required",
}


def parse_header(text):
    """'SFW 58.9kg' -> ('SFW', '58.9'); 'HW +90.7kg' -> ('HW', '+90.7'); else None."""
    if not text:
        return None
    m = HEADER_RE.match(text.strip())
    if not m:
        return None
    return m.group(1).upper(), m.group(2)


def parse_name(text):
    """Returns (name, requirement, flag) or None for blank/'Vacant' cells."""
    if text is None:
        return None
    s = str(text).strip()
    if not s or s.lower() == "vacant":
        return None

    flag = None
    m = WORLD_CHAMP_RE.search(s)
    if m:
        flag = "World Champion"
        s = s[: m.start()].strip()

    requirement = None
    m = FOOTNOTE_RE.search(s)
    if m:
        stars = len(m.group(0).strip())
        requirement = REQ_BY_STARS.get(stars)
        s = s[: m.start()].strip()

    name = re.sub(r"\s+", " ", s).strip()
    if not name:
        return None
    return name, requirement, flag


def get_or_create_fighter(db, cache, name, gender):
    key = (name, gender)
    if key not in cache:
        f = (
            db.query(models.Fighter)
            .filter_by(name=name, gender=models.Gender(gender))
            .one_or_none()
        )
        if f is None:
            f = models.Fighter(name=name, gender=models.Gender(gender))
            db.add(f)
            db.flush()
        cache[key] = f
    return cache[key]


def get_or_create_weight_class(db, cache, gender, code, kg):
    key = (gender, code)
    if key not in cache:
        wc = (
            db.query(models.WeightClass)
            .filter_by(gender=models.Gender(gender), code=code)
            .one_or_none()
        )
        if wc is None:
            wc = models.WeightClass(
                gender=models.Gender(gender),
                code=code,
                name=NAME_BY_CODE.get(code, code),
                kg_display=f"{kg}kg",
                sort_order=0,
            )
            db.add(wc)
            db.flush()
        cache[key] = wc
    return cache[key]


def normalize_scope(scope):
    """Translate spreadsheet shorthand like 'NZ' into the database enum value."""
    raw = (scope or "New Zealand").strip()
    mapping = {
        "NZ": "New Zealand",
        "N.Z.": "New Zealand",
        "New Zealand": "New Zealand",
        "Oceania": "Oceania",
        "International": "International",
        "World": "World",
    }
    normalized = mapping.get(raw, raw)
    try:
        return models.TitleScope(normalized).value
    except ValueError:
        return models.TitleScope.NZ.value


def get_or_create_title(db, cache, wc, level, scope="New Zealand"):
    scope_value = normalize_scope(scope)
    key = (wc.id, level, scope_value)
    if key not in cache:
        t = (
            db.query(models.Title)
            .filter_by(
                weight_class_id=wc.id,
                level=models.TitleLevel(level).value,
                scope=scope_value,
            )
            .one_or_none()
        )
        if t is None:
            t = models.Title(
                weight_class_id=wc.id,
                level=models.TitleLevel(level).value,
                scope=scope_value,
            )
            db.add(t)
            db.flush()
        cache[key] = t
    return cache[key]


def apply_board(db, caches, ws, gender, level, header_row, champion_row, last_row, col_range, as_of):
    fighters, wclasses, titles = caches
    cols = []
    for c in col_range:
        parsed = parse_header(ws.cell(row=header_row, column=c).value)
        if parsed:
            code, kg = parsed
            cols.append((c, code, kg))

    stats = {"champions_set": 0, "champions_skipped": 0, "rankings": 0, "divisions": len(cols)}

    for c, code, kg in cols:
        wc = get_or_create_weight_class(db, wclasses, gender, code, kg)
        title = get_or_create_title(db, titles, wc, level)

        champ = parse_name(ws.cell(row=champion_row, column=c).value)
        if champ:
            name, req, flag = champ
            fighter = get_or_create_fighter(db, fighters, name, gender)
            existing = db.query(models.TitleFight).filter_by(title_id=title.id).first()
            if existing is None:
                db.add(
                    models.TitleFight(
                        title_id=title.id,
                        fight_date=as_of,
                        winner_id=fighter.id,
                        opponent_id=None,
                        event=None,
                        is_vacant_win=False,
                        note="Champion imported from WBC rankings spreadsheet; exact title-win history not recorded.",
                    )
                )
                stats["champions_set"] += 1
            else:
                stats["champions_skipped"] += 1
                print(
                    f"  [skip] {gender}/{level}/{code}: sheet says champion is "
                    f"'{name}' but title already has fight history — not overwritten."
                )

        # Rankings are replaced wholesale from the sheet on every run.
        db.query(models.Ranking).filter_by(title_id=title.id).delete()
        rank = 0
        seen_fighter_ids = set()
        for r in range(champion_row + 1, last_row + 1):
            rank += 1
            parsed = parse_name(ws.cell(row=r, column=c).value)
            if not parsed:
                continue
            name, req, flag = parsed
            fighter = get_or_create_fighter(db, fighters, name, gender)
            if fighter.id in seen_fighter_ids:
                stats["duplicate_entries"] = stats.get("duplicate_entries", 0) + 1
                print(
                    f"  [dup]  {gender}/{level}/{code}: '{name}' appears more than once "
                    f"in this division (rank #{rank} skipped) — check the source sheet."
                )
                continue
            seen_fighter_ids.add(fighter.id)
            db.add(
                models.Ranking(
                    title_id=title.id,
                    fighter_id=fighter.id,
                    rank=rank,
                    requirement=req,
                    flag=flag,
                )
            )
            stats["rankings"] += 1

    return stats


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    Base.metadata.create_all(bind=engine)  # no-op if db/schema.sql already ran

    db = SessionLocal()
    caches = ({}, {}, {})  # fighters, weight_classes, titles
    as_of = date.today()

    try:
        pro_ws = wb["PRO"]
        am_ws = wb["AMATEUR"]
        cons_ws = wb["Consolidated"]

        results = {
            "PRO men (PRO sheet)": apply_board(
                db, caches, pro_ws, "M", "Pro",
                header_row=3, champion_row=4, last_row=14, col_range=range(2, 12), as_of=as_of,
            ),
            "AMATEUR men (AMATEUR sheet)": apply_board(
                db, caches, am_ws, "M", "Amateur",
                header_row=3, champion_row=4, last_row=14, col_range=range(2, 12), as_of=as_of,
            ),
            "Women — Pro (Consolidated sheet)": apply_board(
                db, caches, cons_ws, "F", "Pro",
                header_row=19, champion_row=20, last_row=28, col_range=[2, 4, 5, 6, 7, 8, 9], as_of=as_of,
            ),
            "Women — Amateur (Consolidated sheet)": apply_board(
                db, caches, cons_ws, "F", "Amateur",
                header_row=19, champion_row=20, last_row=28, col_range=[2, 4, 5, 6, 7, 8, 9], as_of=as_of,
            ),
        }

        db.commit()

        fighters, wclasses, titles = caches
        print(
            f"\nImported: {len(fighters)} fighters, {len(wclasses)} weight classes, "
            f"{len(titles)} titles touched.\n"
        )
        for label, s in results.items():
            print(
                f"  {label}: {s['divisions']} division(s), "
                f"{s['champions_set']} champion(s) set, "
                f"{s['champions_skipped']} champion(s) skipped, "
                f"{s['rankings']} ranking row(s)."
            )
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    default_path = "wbc_rankings.xlsx"
    main(sys.argv[1] if len(sys.argv) > 1 else default_path)
